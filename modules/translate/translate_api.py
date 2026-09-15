import os
import io
import time
import json
import uuid
from pathlib import Path
from flask import Blueprint, request, jsonify, send_file, current_app

import docx
from docx.text.paragraph import Paragraph
from docx.table import Table

try:
    from deep_translator import GoogleTranslator, MyMemoryTranslator
    DEEP_TRANSLATOR_AVAILABLE = True
except ImportError:
    DEEP_TRANSLATOR_AVAILABLE = False

from modules.office.office_api import (
    convert_legacy_to_modern_bytes,
    convert_modern_to_legacy_bytes,
    LEGACY_WORD_EXTENSIONS,
    OLE_MAGIC
)

translate_bp = Blueprint('translate_bp', __name__)

LANGUAGE_CODES = {
    'auto': 'auto',
    'vi': 'vi',
    'en': 'en',
    'zh-CN': 'zh-CN',
    'zh-TW': 'zh-TW',
    'ja': 'ja',
    'ko': 'ko',
    'fr': 'fr',
    'de': 'de',
    'ru': 'ru',
    'es': 'es',
    'th': 'th',
    'lo': 'lo',
    'km': 'km'
}


def _get_document_paragraphs(doc: docx.Document):
    """
    Duyệt toàn bộ các đoạn văn bản trong file Word theo đúng thứ tự tài liệu:
    1. Header của các Section (khử trùng lặp node XML và header lặp lại giữa các section)
    2. Toàn bộ Body (Paragraphs và các ô trong Tables, khử trùng lặp do merged cells)
    3. Footer của các Section (khử trùng lặp node XML và footer lặp lại giữa các section)
    Trả về danh sách các đối tượng Paragraph độc nhất (không trùng lặp node XML).
    """
    elements = []
    seen_elements = set()
    seen_hf_text = set()

    def _add_p(p, is_hf=False):
        if not p:
            return
        txt = p.text.strip()
        if not txt:
            return
        if p._p in seen_elements:
            return
        if is_hf:
            norm_hf = " ".join(txt.split())
            if norm_hf in seen_hf_text:
                return
            seen_hf_text.add(norm_hf)
        seen_elements.add(p._p)
        elements.append(p)

    # 1. Header
    for section in doc.sections:
        for p in section.header.paragraphs:
            _add_p(p, is_hf=True)
        for tbl in section.header.tables:
            for row in tbl.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        _add_p(p, is_hf=True)

    # 2. Body (Theo đúng thứ tự cây XML body)
    for child in doc.element.body:
        if child.tag.endswith('p'):
            p = Paragraph(child, doc)
            _add_p(p, is_hf=False)
        elif child.tag.endswith('tbl'):
            t = Table(child, doc)
            for row in t.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        _add_p(p, is_hf=False)

    # 3. Footer
    for section in doc.sections:
        for p in section.footer.paragraphs:
            _add_p(p, is_hf=True)
        for tbl in section.footer.tables:
            for row in tbl.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        _add_p(p, is_hf=True)

    return elements


def replace_paragraph_text(p: Paragraph, new_text: str):
    """
    Thay thế nội dung đoạn văn bằng văn bản mới mà vẫn giữ nguyên:
    - Phong cách (Style), cỡ chữ, font chữ, màu sắc, in đậm/nghiêng của Run đầu tiên.
    - Căn lề (căn giữa, căn trái, căn phải, giãn dòng) của Paragraph.
    """
    if not p.runs:
        p.text = new_text
        return

    p.runs[0].text = new_text
    for r in p.runs[1:]:
        r.text = ''


import re

def should_translate(text: str) -> bool:
    """
    Kiểm tra xem chuỗi văn bản có cần dịch hay không.
    Các chuỗi chỉ gồm số, ngày tháng, công thức, ký tự đặc biệt, đơn vị đo lường,
    hoặc không chứa chữ cái sẽ KHÔNG cần dịch và được giữ nguyên 100%.
    """
    if not text:
        return False
    s = text.strip()
    if not s:
        return False

    # 1. Nếu hoàn toàn không chứa bất kỳ chữ cái nào (chỉ có số, ký tự đặc biệt, dấu câu, toán học, ngày tháng)
    if not any(c.isalpha() for c in s):
        return False

    # 2. URL, email, đường dẫn tập tin
    if re.fullmatch(r'(https?://\S+|www\.\S+|[\w\.-]+@[\w\.-]+\.\w+|[a-zA-Z]:\\[\S\\]+)', s):
        return False

    # 3. Số kèm đơn vị đo lường thông dụng (100km/h, 50kg, $100, 25%, 50Hz, 20°C, v.v.)
    unit_pattern = (
        r'[\$€£¥₫]?\s*[\d\.,\s\+\-\±\~/\:]+\s*'
        r'(km/h|m/s|km|m|cm|mm|kg|g|mg|l|ml|v|w|kw|kwh|hz|khz|mhz|ghz|a|ma|pa|kpa|bar|psi|rpm|fps|kb|mb|gb|tb|%|°c|°f|px|pt|em|rem|m2|m3|cm2)?\s*'
    )
    if re.fullmatch(unit_pattern, s, re.IGNORECASE):
        return False

    # 4. Số La Mã hoặc chỉ mục danh sách đơn lẻ (I, II, III, IV, A., b), 1.1, v.v.)
    if re.fullmatch(r'[IVXLCDMivxlcdm]+[\.\)\:\-]?\s*', s) or re.fullmatch(r'\(?[a-zA-Z]\)?[\.\:\-]?\s*', s):
        return False

    return True


def translate_sentences_google(texts: list[str], src_lang: str = 'auto', tgt_lang: str = 'vi') -> dict[str, str]:
    """
    Dịch danh sách các câu bằng Google Translate với kỹ thuật khử trùng lặp và phân nhóm thông minh:
    - Bỏ qua các chuỗi chỉ có số, ký tự đặc biệt, đơn vị đo lường (giữ nguyên gốc).
    - Khử trùng lặp để tiết kiệm tối đa lượt gọi API.
    - Ghép nối các câu ngắn bằng dấu xuống dòng để dịch hàng loạt trong 1 request.
    - Tự động fallback sang từng câu riêng lẻ nếu gộp nhóm không khớp.
    """
    if not DEEP_TRANSLATOR_AVAILABLE:
        raise RuntimeError("Thư viện deep-translator chưa được cài đặt.")

    src = LANGUAGE_CODES.get(src_lang, 'auto')
    tgt = LANGUAGE_CODES.get(tgt_lang, 'vi')

    # Phân loại: chuỗi cần dịch và chuỗi giữ nguyên (số, ký hiệu, đơn vị)
    unique_texts = []
    seen = set()
    translation_map = {}

    for t in texts:
        clean = t.strip()
        if not clean or clean in seen:
            continue
        seen.add(clean)

        # Nếu là số hoặc ký hiệu -> Giữ nguyên, không dịch
        if not should_translate(clean):
            translation_map[clean] = clean
        else:
            unique_texts.append(clean)

    if not unique_texts:
        return translation_map

    # Chia nhóm để dịch gộp (mỗi nhóm tối đa 15 câu và không quá 1500 ký tự)
    chunks = []
    current_chunk = []
    current_len = 0

    for t in unique_texts:
        if len(current_chunk) >= 15 or (current_len + len(t)) > 1500:
            chunks.append(current_chunk)
            current_chunk = [t]
            current_len = len(t)
        else:
            current_chunk.append(t)
            current_len += len(t) + 1
    if current_chunk:
        chunks.append(current_chunk)

    for chunk in chunks:
        combined = '\n'.join(chunk)
        success = False

        try:
            res = translator.translate(combined)
            res_lines = res.split('\n')
            # Kiểm tra nếu số dòng trả về khớp với số câu trong nhóm
            if len(res_lines) == len(chunk):
                for orig, trans in zip(chunk, res_lines):
                    translation_map[orig] = trans.strip()
                success = True
            time.sleep(0.05)
        except Exception:
            pass

        # Nếu dịch gộp thất bại (do số dòng lệch hoặc lỗi cú pháp), dịch từng câu với độ trễ an toàn
        if not success:
            for t in chunk:
                try:
                    res_single = translator.translate(t)
                    translation_map[t] = res_single.strip()
                    time.sleep(0.08)
                except Exception:
                    # Nếu Google chặn tạm thời, thử fallback sang MyMemory
                    try:
                        mymemory_src = 'en-US' if src == 'en' else src
                        mymemory_tgt = 'vi-VN' if tgt == 'vi' else tgt
                        mm_trans = MyMemoryTranslator(source=mymemory_src, target=mymemory_tgt)
                        translation_map[t] = mm_trans.translate(t).strip()
                    except Exception:
                        translation_map[t] = t  # Giữ nguyên nếu tất cả đều lỗi

    return translation_map


# ── API ENDPOINTS ────────────────────────────────────────────────────────────

@translate_bp.route('/extract', methods=['POST'])
def extract_segments():
    """
    Trích xuất toàn bộ các đoạn văn bản từ file Word phục vụ chế độ Tự dịch song ngữ.
    """
    if 'file' not in request.files:
        return jsonify({'error': 'Không tìm thấy tệp tin trong yêu cầu'}), 400

    file = request.files['file']
    if not file.filename:
        return jsonify({'error': 'Tên tệp không hợp lệ'}), 400

    orig_name = file.filename
    ext = os.path.splitext(orig_name)[1].lower()

    if ext not in {'.docx', '.doc', '.docm'}:
        return jsonify({'error': 'Chỉ hỗ trợ tệp tài liệu Word (.docx, .doc)'}), 400

    data = file.read()

    # Xử lý nếu là file .doc cũ
    if ext in LEGACY_WORD_EXTENSIONS or data.startswith(OLE_MAGIC):
        try:
            data, _ = convert_legacy_to_modern_bytes(data, ext)
        except Exception as e:
            return jsonify({'error': f'Lỗi chuyển đổi file .doc cũ sang docx: {str(e)}'}), 400

    try:
        doc = docx.Document(io.BytesIO(data))
        paragraphs = _get_document_paragraphs(doc)

        from collections import Counter
        text_counts = Counter(p.text.strip() for p in paragraphs)

        segments = []
        for idx, p in enumerate(paragraphs):
            raw_text = p.text.strip()
            is_trans = should_translate(raw_text)
            segments.append({
                'id': idx,
                'text': raw_text,
                'count': text_counts[raw_text],
                'translatable': is_trans,
                'translated': raw_text if not is_trans else ''
            })

        return jsonify({
            'filename': orig_name,
            'total_segments': len(segments),
            'segments': segments
        })

    except Exception as e:
        return jsonify({'error': f'Lỗi khi đọc file Word: {str(e)}'}), 500


@translate_bp.route('/suggest', methods=['POST'])
def suggest_translation():
    """
    Gợi ý bản dịch nhanh cho danh sách câu trong bảng biên tập song ngữ.
    """
    body = request.get_json(silent=True) or {}
    texts = body.get('texts', [])
    src_lang = body.get('src_lang', 'auto')
    tgt_lang = body.get('tgt_lang', 'vi')

    if not texts:
        return jsonify({'error': 'Danh sách văn bản trống'}), 400

    try:
        translation_map = translate_sentences_google(texts, src_lang=src_lang, tgt_lang=tgt_lang)
        results = [translation_map.get(t.strip(), t) for t in texts]
        return jsonify({'translations': results})
    except Exception as e:
        return jsonify({'error': f'Lỗi dịch thuật: {str(e)}'}), 500


@translate_bp.route('/auto', methods=['POST'])
def translate_auto():
    """
    Dịch tự động toàn bộ file Word qua Google Dịch và trả về file Word hoàn chỉnh.
    """
    if 'file' not in request.files:
        return jsonify({'error': 'Không tìm thấy tệp tin trong yêu cầu'}), 400

    file = request.files['file']
    if not file.filename:
        return jsonify({'error': 'Tên tệp không hợp lệ'}), 400

    src_lang = request.form.get('src_lang', 'auto')
    tgt_lang = request.form.get('tgt_lang', 'vi')

    orig_name = file.filename
    base_name, ext = os.path.splitext(orig_name)
    ext = ext.lower()

    if ext not in {'.docx', '.doc', '.docm'}:
        return jsonify({'error': 'Chỉ hỗ trợ tệp tài liệu Word (.docx, .doc)'}), 400

    data = file.read()

    # Nếu là file .doc cũ, chuyển sang docx
    is_legacy = ext in LEGACY_WORD_EXTENSIONS or data.startswith(OLE_MAGIC)
    if is_legacy:
        try:
            data, _ = convert_legacy_to_modern_bytes(data, ext)
        except Exception as e:
            return jsonify({'error': f'Lỗi chuyển đổi file .doc cũ sang docx: {str(e)}'}), 400

    try:
        doc = docx.Document(io.BytesIO(data))
        paragraphs = _get_document_paragraphs(doc)

        if not paragraphs:
            return jsonify({'error': 'Tài liệu không chứa văn bản nào để dịch'}), 400

        all_texts = [p.text.strip() for p in paragraphs]
        translation_map = translate_sentences_google(all_texts, src_lang=src_lang, tgt_lang=tgt_lang)

        # Bơm văn bản đã dịch vào các Paragraph tương ứng
        norm_trans_map = {" ".join(k.split()): v for k, v in translation_map.items()}
        def _get_auto_trans(txt):
            clean = txt.strip()
            if clean in translation_map:
                return translation_map[clean]
            norm = " ".join(clean.split())
            if norm in norm_trans_map:
                return norm_trans_map[norm]
            return None

        for p in paragraphs:
            trans_val = _get_auto_trans(p.text)
            if trans_val:
                replace_paragraph_text(p, trans_val)

        # Đồng bộ toàn bộ Header và Footer trên mọi section có cùng nội dung gốc
        for s in doc.sections:
            for p in s.header.paragraphs:
                trans_val = _get_auto_trans(p.text)
                if trans_val:
                    replace_paragraph_text(p, trans_val)
            for tbl in s.header.tables:
                for row in tbl.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            trans_val = _get_auto_trans(p.text)
                            if trans_val:
                                replace_paragraph_text(p, trans_val)
            for p in s.footer.paragraphs:
                trans_val = _get_auto_trans(p.text)
                if trans_val:
                    replace_paragraph_text(p, trans_val)
            for tbl in s.footer.tables:
                for row in tbl.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            trans_val = _get_auto_trans(p.text)
                            if trans_val:
                                replace_paragraph_text(p, trans_val)

        out_buffer = io.BytesIO()
        doc.save(out_buffer)
        out_bytes = out_buffer.getvalue()

        download_name = f"{base_name}_translated.docx"
        mimetype = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'

        # Nếu file gốc là .doc và người dùng muốn giữ .doc
        if is_legacy and ext == '.doc':
            try:
                out_bytes = convert_modern_to_legacy_bytes(out_bytes, '.doc')
                download_name = f"{base_name}_translated.doc"
                mimetype = 'application/msword'
            except Exception:
                pass

        resp = send_file(
            io.BytesIO(out_bytes),
            mimetype=mimetype,
            as_attachment=True,
            download_name=download_name
        )
        resp.headers['X-Translate-Stats'] = json.dumps({
            'total_segments': len(paragraphs),
            'unique_segments': len(translation_map)
        })
        resp.headers['Access-Control-Expose-Headers'] = 'X-Translate-Stats'
        return resp

    except Exception as e:
        return jsonify({'error': f'Lỗi trong quá trình dịch thuật: {str(e)}'}), 500


@translate_bp.route('/apply', methods=['POST'])
def apply_translations():
    """
    Bơm danh sách câu đã dịch do người dùng tự biên tập vào file Word gốc.
    """
    if 'file' not in request.files:
        return jsonify({'error': 'Không tìm thấy tệp Word gốc'}), 400

    file = request.files['file']
    translations_raw = request.form.get('translations', '')

    if not translations_raw:
        return jsonify({'error': 'Thiếu dữ liệu bản dịch'}), 400

    try:
        translations = json.loads(translations_raw)
    except Exception:
        return jsonify({'error': 'Định dạng dữ liệu bản dịch không hợp lệ'}), 400

    orig_name = file.filename
    base_name, ext = os.path.splitext(orig_name)
    ext = ext.lower()

    data = file.read()
    is_legacy = ext in LEGACY_WORD_EXTENSIONS or data.startswith(OLE_MAGIC)
    if is_legacy:
        try:
            data, _ = convert_legacy_to_modern_bytes(data, ext)
        except Exception as e:
            return jsonify({'error': f'Lỗi chuyển đổi file .doc cũ sang docx: {str(e)}'}), 400

    try:
        doc = docx.Document(io.BytesIO(data))
        paragraphs = _get_document_paragraphs(doc)

        # translations là mảng [{ id: 0, text: '...', orig: '...' }, ...] hoặc dict { "0": "..." }
        trans_dict = {}
        text_map = {}
        norm_text_map = {}
        if isinstance(translations, list):
            for item in translations:
                if 'id' in item:
                    trans_dict[int(item['id'])] = item.get('text', '')
                if 'orig' in item and item.get('text', '').strip():
                    orig_clean = item['orig'].strip()
                    val = item.get('text', '').strip()
                    text_map[orig_clean] = val
                    norm_text_map[" ".join(orig_clean.split())] = val
        elif isinstance(translations, dict):
            for k, v in translations.items():
                trans_dict[int(k)] = v

        def _get_matched_trans(txt):
            clean = txt.strip()
            if clean in text_map:
                return text_map[clean]
            norm = " ".join(clean.split())
            if norm in norm_text_map:
                return norm_text_map[norm]
            return None

        for idx, p in enumerate(paragraphs):
            matched = _get_matched_trans(p.text)
            if idx in trans_dict and trans_dict[idx].strip():
                replace_paragraph_text(p, trans_dict[idx])
            elif matched:
                replace_paragraph_text(p, matched)

        # Đồng bộ toàn bộ Header và Footer trên mọi section có cùng nội dung gốc
        if text_map or norm_text_map:
            for s in doc.sections:
                for p in s.header.paragraphs:
                    val = _get_matched_trans(p.text)
                    if val:
                        replace_paragraph_text(p, val)
                for tbl in s.header.tables:
                    for row in tbl.rows:
                        for cell in row.cells:
                            for p in cell.paragraphs:
                                val = _get_matched_trans(p.text)
                                if val:
                                    replace_paragraph_text(p, val)
                for p in s.footer.paragraphs:
                    val = _get_matched_trans(p.text)
                    if val:
                        replace_paragraph_text(p, val)
                for tbl in s.footer.tables:
                    for row in tbl.rows:
                        for cell in row.cells:
                            for p in cell.paragraphs:
                                val = _get_matched_trans(p.text)
                                if val:
                                    replace_paragraph_text(p, val)

        out_buffer = io.BytesIO()
        doc.save(out_buffer)
        out_bytes = out_buffer.getvalue()

        download_name = f"{base_name}_translated.docx"
        mimetype = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'

        if is_legacy and ext == '.doc':
            try:
                out_bytes = convert_modern_to_legacy_bytes(out_bytes, '.doc')
                download_name = f"{base_name}_translated.doc"
                mimetype = 'application/msword'
            except Exception:
                pass

        resp = send_file(
            io.BytesIO(out_bytes),
            mimetype=mimetype,
            as_attachment=True,
            download_name=download_name
        )
        return resp

    except Exception as e:
        return jsonify({'error': f'Lỗi khi áp dụng bản dịch vào file Word: {str(e)}'}), 500
